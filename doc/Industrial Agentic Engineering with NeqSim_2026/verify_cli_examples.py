"""Exercise the real NeqSim CLI with fixture-only defaults and documents.

The child-process startup hook redirects only NeqSim's task-defaults pathname.
It leaves HOME and USERPROFILE unchanged. Catalog listing additionally points
private catalog/installation metadata at empty fixture paths; public repository
catalogs still come from the selected source checkout. No install command runs.
"""
from book_runtime import BOOK, LOCAL_DEPS
from pathlib import Path
import argparse
import hashlib
import json
import os
import re
import subprocess
import shutil
import sys
import tempfile
import time
from datetime import datetime, timezone
import yaml
from docx import Document
from docx.shared import Pt
from pypdf import PdfWriter


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else None


def run(project_root):
    source = Path(project_root).resolve()
    cli = source / "devtools/neqsim_cli.py"
    verification = BOOK / "verification"
    fixtures_parent = verification / "cli_fixtures"
    fixtures_parent.mkdir(exist_ok=True)
    fixture = Path(tempfile.mkdtemp(prefix="2026-09-13_", dir=str(fixtures_parent)))
    hook = fixture / "python startup"
    hook.mkdir()
    defaults = fixture / "isolated settings/task_defaults.json"
    defaults.parent.mkdir()
    user_defaults = Path.home() / ".neqsim/task_defaults.json"
    before = digest(user_defaults)
    source_commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=str(source), text=True).strip()
    records = []
    assertions = []
    harness_error = None
    start = time.monotonic()
    hook_code = r'''import os
import sys
from pathlib import Path
_original_expanduser = os.path.expanduser
_fixture_defaults = os.environ.get("NEQSIM_BOOK_TEST_DEFAULTS")
if _fixture_defaults:
    def _isolated_expanduser(path):
        if os.fspath(path).replace("\\", "/") == "~/.neqsim/task_defaults.json":
            return _fixture_defaults
        return _original_expanduser(path)
    os.path.expanduser = _isolated_expanduser
if len(sys.argv) > 1 and sys.argv[1] in ("agent", "skill"):
    sys.path.insert(0, os.environ["NEQSIM_BOOK_TEST_DEVTOOLS"])
    import install_skill
    import install_agent
    _catalog_root = Path(os.environ["NEQSIM_BOOK_TEST_DEFAULTS"]).parent
    for _module, _kind in ((install_skill, "skills"), (install_agent, "agents")):
        _module.PRIVATE_CATALOG_FILE = _catalog_root / ("private-" + _kind + ".yaml")
        _module.INSTALL_DIR = _catalog_root / _kind
        _module.MANIFEST_FILE = _module.INSTALL_DIR / "installed.json"
        _module.EXPORT_DIR = _catalog_root / "export"
        _module.EXPORT_PROFILE_FILE = _module.EXPORT_DIR / "export-profile.json"
'''
    (hook / "sitecustomize.py").write_text(hook_code, encoding="utf-8")
    child_env = dict(os.environ)
    for name in ("NEQSIM_TASK_ROOT", "NEQSIM_DOCUMENT_ROOT", "NEQSIM_REPORT_TEMPLATE", "NEQSIM_TASK_DIR", "NEQSIM_REPORT_TITLE", "NEQSIM_REPORT_AUTHOR"):
        child_env.pop(name, None)
    child_env.update(NEQSIM_PROJECT_ROOT=str(source), NEQSIM_BOOK_TEST_DEFAULTS=str(defaults),
                     NEQSIM_BOOK_TEST_DEVTOOLS=str(source / "devtools"), PYTHONUTF8="1",
                     PYTHONPATH=os.pathsep.join([str(hook), str(LOCAL_DEPS), str(BOOK)]))

    def check(name, argv, expected=0, contains=(), absent=(), overrides=None, timeout=240, cwd=None, entry=None):
        env = dict(child_env)
        env.update(overrides or {})
        command = [sys.executable, str(entry or cli)] + [str(a) for a in argv]
        try:
            completed = subprocess.run(command, cwd=str(cwd or fixture), env=env,
                                       input="", capture_output=True, text=True, encoding="utf-8",
                                       errors="replace", timeout=timeout)
            output = completed.stdout + completed.stderr
            code = completed.returncode
        except subprocess.TimeoutExpired as error:
            output, code = "Command timed out after {} seconds: {}".format(timeout, error), None
        log = fixture / ("{:02d}_".format(len(records) + 1) + name + ".log")
        log.write_text(output, encoding="utf-8")
        codes = expected if isinstance(expected, tuple) else (expected,)
        failures = ([] if code in codes else ["Exit code {} was not in {}".format(code, codes)])
        failures += ["Missing expected text: " + text for text in contains if text not in output]
        failures += ["Unexpected text: " + text for text in absent if text in output]
        records.append({"name": name, "command": command, "cwd": str(cwd or fixture),
                        "environment_overrides": overrides or {}, "exit_code": code,
                        "expected_exit_codes": list(codes), "status": "passed" if not failures else "failed",
                        "failures": failures, "log": log.relative_to(BOOK).as_posix()})
        print(name + ": " + records[-1]["status"], flush=True)
        return output

    def require(name, condition, detail):
        assertions.append({"name": name, "status": "passed" if condition else "failed", "detail": detail})
        if not condition:
            print(name + ": failed", flush=True)

    def make_template(path, header):
        document = Document()
        document.styles["Normal"].font.name = "Garamond"
        document.styles["Normal"].font.size = Pt(12)
        document.sections[0].header.paragraphs[0].text = header
        document.add_paragraph("FIXTURE TEMPLATE BOILERPLATE")
        document.save(str(path))

    try:
        # Prove that child startup is active before executing any setting writes.
        probe = subprocess.run([sys.executable, "-c", "import os; print(os.path.expanduser('~/.neqsim/task_defaults.json'))"],
                               env=child_env, capture_output=True, text=True, encoding="utf-8", check=True)
        if Path(probe.stdout.strip()).resolve() != defaults.resolve():
            raise RuntimeError("Default-path isolation was not active; refused to run setting commands")
        check("help", ["--help"], contains=("new-task", "documents", "work-record", "report"))
        check("onboard_help", ["onboard", "--help"], contains=("--check",))
        check("try_help", ["try", "--help"])
        check("show_unset_document_root", ["--show-document-root"], contains=("none",))
        check("documents_unset", ["documents"], expected=2, contains=("No document root configured",))
        check("show_unset_template", ["--show-report-template"], contains=("built-in styling",))

        task_root = fixture / "Saved task destination"
        env_task_root = fixture / "Environment task destination"
        explicit_root = fixture / "Explicit task destination"
        documents = fixture / "Engineering Documents"
        env_documents = fixture / "Override Documents"
        empty_documents = fixture / "Empty Documents"
        for path in (task_root, env_task_root, explicit_root, documents / "vendor data/nested", documents / "standards", documents / ".private", env_documents, empty_documents):
            path.mkdir(parents=True, exist_ok=True)
        (documents / "README.txt").write_text("Public verification fixtures; not engineering evidence.\n", encoding="utf-8")
        (documents / "vendor data/nested/notes.txt").write_text("Nested fixture\n", encoding="utf-8")
        (documents / ".hidden.txt").write_text("Hidden fixture\n", encoding="utf-8")
        (env_documents / "override.txt").write_text("Environment override fixture\n", encoding="utf-8")
        for path in (documents / "vendor data/Compressor Datasheet.PDF", documents / "standards/Hydrate guidance.pdf", documents / ".private/ignored.pdf"):
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            with path.open("wb") as handle:
                writer.write(handle)
        documents_before = {p.relative_to(documents).as_posix(): digest(p) for p in documents.rglob("*") if p.is_file()}
        check("set_task_root", ["--set-task-root", task_root], contains=(str(task_root),))
        check("show_saved_task_root", ["--show-task-root"], contains=(str(task_root),))
        check("task_environment_override", ["--show-task-root"], overrides={"NEQSIM_TASK_ROOT": str(env_task_root)}, contains=(str(env_task_root),))
        check("task_explicit_override", ["new-task", "--show-task-root", "--task-root", explicit_root],
              overrides={"NEQSIM_TASK_ROOT": str(env_task_root)}, contains=(str(explicit_root),))
        check("set_document_root_unquoted_spaces", ["set-document-root"] + str(documents).split(" "), contains=(str(documents),))
        check("show_saved_document_root", ["--show-document-root"], contains=(str(documents),))
        check("documents_recursive", ["documents"], contains=("4 document(s)", "Compressor Datasheet.PDF", "Hydrate guidance.pdf", "notes.txt"), absent=("ignored.pdf", ".hidden.txt"))
        check("documents_case_insensitive", ["documents", "COMPRESSOR"], contains=("1 document(s)", "Compressor Datasheet.PDF"))
        check("documents_pdf_substring", ["documents", ".PDF"], contains=("2 document(s)",))
        check("documents_path_substring", ["documents", "vendor", "data"], contains=("2 document(s)",))
        check("documents_unmatched_multiword_substring", ["documents", "API", "521"], expected=1, contains=("No documents",))
        check("documents_not_glob", ["documents", "*.pdf"], expected=1, contains=("No documents",))
        check("documents_environment_override", ["documents"], overrides={"NEQSIM_DOCUMENT_ROOT": str(env_documents)}, contains=("override.txt",), absent=("Compressor Datasheet",))
        check("documents_empty", ["documents"], expected=1, overrides={"NEQSIM_DOCUMENT_ROOT": str(empty_documents)}, contains=("No documents",))
        check("documents_missing_root", ["documents"], expected=2, overrides={"NEQSIM_DOCUMENT_ROOT": str(fixture / "Missing documents")}, contains=("Document root folder not found",))

        prompt = fixture / "Study request with spaces.txt"
        prompt.write_text("Verify CLI task creation and report metadata using fixture data only.", encoding="utf-8")
        check("new_task_explicit_destination", ["new-task", "Gas compression screening", "--type", "B", "--scale", "standard",
              "--report-depth", "standard", "--intake-pause", "always", "--author", "CLI Verification", "--task-root", explicit_root,
              "--prompt-file", prompt], overrides={"NEQSIM_TASK_ROOT": str(env_task_root)})
        tasks = [p for p in explicit_root.iterdir() if p.is_dir() and p.name != "TASK_TEMPLATE"]
        if len(tasks) != 1:
            raise RuntimeError("Task creation did not produce exactly one isolated task")
        task = tasks[0]
        study = yaml.safe_load((task / "study_config.yaml").read_text(encoding="utf-8"))
        require("new_task_persists_document_root", Path(study["inputs"]["document_root"]).resolve() == documents.resolve(), "inputs.document_root retains the source library.")
        require("new_task_title_and_scale", study["study"]["title"] == "Gas compression screening" and study["study"]["scale"] == "standard", "study identity matches the command.")
        require("new_task_stays_in_explicit_root", task.parent == explicit_root and not list(env_task_root.iterdir()), "Explicit task destination overrides saved and environment paths.")
        check("new_task_document_setting_alias", ["new-task", "--show-document-root"], contains=(str(documents),))
        check("new_task_template_alias_unset", ["new-task", "--show-report-template"], contains=("built-in styling",))
        quick_root = fixture / "Quick task destination"
        check("new_quick_property_task", ["new-task", "Methane density check", "--type", "A", "--scale", "quick",
              "--report-depth", "brief", "--task-root", quick_root])
        quick_tasks = [p for p in quick_root.iterdir() if p.is_dir() and p.name != "TASK_TEMPLATE"]
        quick_config = yaml.safe_load((quick_tasks[0] / "study_config.yaml").read_text(encoding="utf-8"))
        require("quick_task_metadata", len(quick_tasks) == 1 and quick_config["study"]["scale"] == "quick"
                and quick_config["report"]["depth"] == "brief", "The Chapter 1 property task records quick scale and brief report depth.")

        # Complete the fixture's metadata without pretending it is a real engineering study.
        study["study"].update(title="CLI fixture report", scale="quick")
        study["analysis"] = {"engine": "script", "scripts": [{"file": "fixture.py", "purpose": "Count declared fixture documents.", "produces": "step2_analysis/count.json"}]}
        study["notebooks"] = {"required": False, "plan": []}
        study["inputs"].update(documents=[], data_sources=[])
        study["report"] = {"depth": "standard", "formats": ["docx", "html"], "work_record": "required"}
        (task / "study_config.yaml").write_text(yaml.safe_dump(study, sort_keys=False), encoding="utf-8")
        (task / "step1_scope_and_research/task_spec.md").write_text("# CLI fixture report\n\n## Objective\nVerify command dispatch, isolated configuration and document generation. This is not a process-design assessment.\n", encoding="utf-8")
        (task / "step2_analysis/fixture.py").write_text('"""CLI test fixture; no engineering calculation."""\n', encoding="utf-8")
        (task / "step2_analysis/count.json").write_text('{"fixture_count": 4}\n', encoding="utf-8")
        (task / "results.json").write_text(json.dumps({"key_results": {"fixture_count": 4}, "validation": {"fixture_complete": True},
              "approach": "Exercise CLI commands with isolated settings and synthetic fixture documents.", "conclusions": "This fixture tests document generation only.",
              "objective": "Verify isolated command dispatch and title-derived output filenames."}, indent=2), encoding="utf-8")
        references = task / "step1_scope_and_research/references"
        references.mkdir(exist_ok=True)
        shutil.copy2(documents / "vendor data/Compressor Datasheet.PDF", references / "vendor_datasheet_fixture.pdf")
        check("organize_fixture_sources", [task, "--organize"], entry=source / "devtools/generate_sources_md.py",
              contains=("SOURCES.md", "collection_manifest.json"))
        manifest = json.loads((references / "collection_manifest.json").read_text(encoding="utf-8"))
        require("source_manifest_has_fixture", "vendor_datasheet_fixture.pdf" in json.dumps(manifest)
                and (references / "SOURCES.md").is_file(), "The real source organizer indexes the copied fixture without altering its source library.")
        check("skill_search", ["gas compression screening", "--top", "5"], entry=source / "devtools/skill_search.py", contains=("Top 5 skills",))
        check("agent_search", ["gas compression screening", "--top", "8"], entry=source / "devtools/agent_search.py", contains=("Top 8 agents",))
        check("consistency_checker_script_fixture", [task], entry=source / "devtools/consistency_checker.py", contains=("CONSISTENCY CHECK SUMMARY", "Status: PASS"))
        consistency = json.loads((task / "consistency_report.json").read_text(encoding="utf-8"))
        require("consistency_gate_scope_recorded", consistency["summary"]["critical"] == 0
                and consistency["summary"]["notebooks_checked"] == 0,
                "Script-fixture invocation succeeds with zero notebooks; this is command verification, not numerical consistency validation.")
        check("work_record_generate", ["work-record", task], contains=("Wrote", "WORK_RECORD.md"))
        check("work_record_unfilled_fails", ["work-record", task, "--check"], expected=1, contains=("template text",))
        work_record = task / "step3_report/WORK_RECORD.md"
        narratives = {"background": "This isolated fixture checks command-line documentation; its files are synthetic.",
                      "method": "Ran real NeqSim CLI commands with a redirected task-defaults file and inspected generated artifacts.",
                      "limitations": "No engineering prediction, field data or physical-validation claim is represented by this fixture."}
        text = work_record.read_text(encoding="utf-8")
        text = re.sub(r"(<!--\s*WORK_RECORD:NARRATIVE id=([A-Za-z0-9_-]+)\s*-->).*?(<!--\s*/WORK_RECORD:NARRATIVE\s*-->)",
                      lambda m: m.group(1) + "\n" + narratives.get(m.group(2), "Fixture-specific narrative reviewed for command verification.") + "\n" + m.group(3), text, flags=re.S)
        work_record.write_text(text, encoding="utf-8")
        check("work_record_filled_passes", ["work-record", task, "--check"], contains=("Work record OK",))
        check("work_record_regeneration_preserves", ["work-record", task], contains=("Preserved narrative",))
        require("work_record_narrative_preserved", all(value in work_record.read_text(encoding="utf-8") for value in narratives.values()), "All three handwritten fixture narratives survive regeneration.")

        saved_template = fixture / "Saved company template.docx"
        explicit_template = fixture / "Explicit company template.docx"
        make_template(saved_template, "SAVED TEMPLATE FIXTURE")
        make_template(explicit_template, "EXPLICIT TEMPLATE FIXTURE")
        check("set_report_template", ["--set-report-template", saved_template], contains=(str(saved_template),))
        check("show_report_template", ["--show-report-template"], contains=(str(saved_template),))
        check("report_saved_template", ["report", task], contains=("CLI_fixture_report.docx", "CLI_fixture_report.html"))
        output_dir = task / "step3_report"
        docx_path = output_dir / "CLI_fixture_report.docx"
        generated = Document(str(docx_path))
        require("report_template_style_inherited", generated.styles["Normal"].font.name == "Garamond", "Normal font inherited from the real DOCX fixture template.")
        require("report_saved_header_inherited", "SAVED TEMPLATE FIXTURE" in " ".join(p.text for s in generated.sections for p in s.header.paragraphs), "The saved template header is preserved.")
        require("report_template_body_removed", "FIXTURE TEMPLATE BOILERPLATE" not in " ".join(p.text for p in generated.paragraphs), "Template body is removed by default.")
        require("report_no_unrequested_paper", not list(output_dir.glob("*_Paper.*")), "Scientific-paper output requires an explicit paper option.")

        check("report_explicit_title_template_override", ["report", task, "--title", "Gas compression screening", "--template", explicit_template],
              overrides={"NEQSIM_REPORT_TEMPLATE": str(saved_template), "NEQSIM_REPORT_TITLE": "Environment title"}, contains=("Gas_compression_screening.docx",))
        renamed = Document(str(output_dir / "Gas_compression_screening.docx"))
        require("report_explicit_header_inherited", "EXPLICIT TEMPLATE FIXTURE" in " ".join(p.text for s in renamed.sections for p in s.header.paragraphs), "Explicit --template wins over environment and saved setting.")
        require("report_title_outputs_exist", (output_dir / "Gas_compression_screening.html").is_file(), "Both output filenames follow the explicit title.")
        require("report_superseded_outputs_removed", not docx_path.exists() and not (output_dir / "CLI_fixture_report.html").exists(), "Only superseded fixture reports are removed on title change.")
        check("report_environment_title_template", ["report", task], overrides={"NEQSIM_REPORT_TITLE": "Environment fixture report", "NEQSIM_REPORT_TEMPLATE": str(explicit_template)}, contains=("Environment_fixture_report.docx",))
        check("report_missing_explicit_template", ["report", task, "--template", fixture / "Missing template.docx"], expected=2, contains=("ERROR:", "Missing template.docx"))
        saved = json.loads(defaults.read_text(encoding="utf-8"))
        saved["report_template"] = str(fixture / "Missing saved template.docx")
        defaults.write_text(json.dumps(saved, indent=2), encoding="utf-8")
        check("show_missing_saved_template", ["--show-report-template"], expected=2, contains=("not usable",))
        check("report_missing_saved_template", ["report", task], expected=2, contains=("ERROR:", "Missing saved template.docx"))
        check("report_explicit_bypasses_missing_saved", ["report", task, "--template", explicit_template], contains=("CLI_fixture_report.docx",))
        check("reset_report_template", ["--reset-report-template"])
        check("reset_document_root", ["--reset-document-root"])
        check("reset_task_root", ["--reset-task-root"])
        check("document_root_undefined_after_reset", ["--show-document-root"], contains=("none",))
        require("resumed_task_keeps_document_root", yaml.safe_load((task / "study_config.yaml").read_text(encoding="utf-8"))["inputs"]["document_root"] == str(documents), "Resetting user defaults does not rewrite an existing task's captured source library.")
        require("document_library_unchanged", documents_before == {p.relative_to(documents).as_posix(): digest(p) for p in documents.rglob("*") if p.is_file()}, "Search and report commands do not mutate source documents.")
        # Diagnostics are observed honestly; missing optional runtime components are retained.
        check("doctor_skip_jar", ["doctor", "--skip-jar"], expected=(0, 1), contains=("Summary:", "Not checked (--skip-jar)"))
        check("doctor_with_jar_check", ["doctor"], expected=(0, 1), contains=("Summary:",))
        check("agent_catalog_list", ["agent", "list"], timeout=300, contains=("Agent Catalog",))
        check("skill_catalog_list", ["skill", "list"], timeout=300, contains=("Skill Catalog",))
    except Exception as error:
        harness_error = {"type": type(error).__name__, "message": str(error)}
        raise
    finally:
        after = digest(user_defaults)
        require("user_defaults_unchanged", before == after, "Actual user task-defaults file is unchanged; HOME and USERPROFILE were never repurposed.")
        report = {"generated_at": datetime.now(timezone.utc).isoformat(), "python": sys.executable,
                  "source_root": str(source), "source_commit": source_commit, "cli_sha256": digest(cli),
                  "verifier_sha256": digest(Path(__file__)), "harness_error": harness_error,
                  "fixture_root": fixture.relative_to(BOOK).as_posix(), "isolation": "Child sitecustomize redirects only the task-defaults path; private catalogs are isolated for list commands.",
                  "user_defaults_unchanged": before == after, "commands": records, "assertions": assertions,
                  "elapsed_seconds": time.monotonic() - start,
                  "scope": "Actual command dispatch, configuration, file search, report artifacts and work-record behavior; not physical validation or a full environment health certificate.",
                  "failures": [r["name"] for r in records + assertions if r["status"] != "passed"] + (["harness_exception"] if harness_error else [])}
        (verification / "cli_examples_2026-09-13.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    if report["failures"]:
        raise RuntimeError("CLI verification failures: " + ", ".join(report["failures"]))
    print(json.dumps({"commands": len(records), "assertions": len(assertions), "failures": report["failures"], "user_defaults_unchanged": before == after}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", required=True)
    run(parser.parse_args().project_root)
